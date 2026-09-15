import base64
import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from backend.ingestion import parse_document


PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _save(deck: Presentation, path: Path) -> Path:
    deck.save(path)
    return path


def test_pptx_skips_image_only_slide_and_extracts_grouped_text(tmp_path: Path):
    deck = Presentation()
    text_slide = deck.slides.add_slide(deck.slide_layouts[6])
    text_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "第一页正文"
    group_slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = group_slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "组合形状中的文字"
    image_slide = deck.slides.add_slide(deck.slide_layouts[6])
    image_slide.shapes.add_picture(io.BytesIO(PNG_1PX), Inches(1), Inches(1), width=Inches(1))
    path = _save(deck, tmp_path / "mixed.pptx")

    chunks = parse_document(path, "doc", "mixed.pptx")

    assert [chunk["locator"]["slide_number"] for chunk in chunks] == [1, 2]
    assert "第一页正文" in chunks[0]["text"]
    assert "组合形状中的文字" in chunks[1]["text"]


def test_pptx_with_only_images_has_clear_error(tmp_path: Path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(PNG_1PX), Inches(1), Inches(1), width=Inches(1))
    path = _save(deck, tmp_path / "image_only.pptx")

    with pytest.raises(ValueError, match="图片型课件"):
        parse_document(path, "doc", "image_only.pptx")


def test_pptx_extracts_native_chart_labels(tmp_path: Path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = ChartData()
    data.categories = ["低收入", "中等收入", "高收入"]
    data.add_series("储蓄率", (10, 20, 30))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "储蓄率对比"
    path = _save(deck, tmp_path / "chart.pptx")

    chunks = parse_document(path, "doc", "chart.pptx")

    assert len(chunks) == 1
    assert "储蓄率对比" in chunks[0]["text"]
    assert "低收入" in chunks[0]["text"]
    assert "储蓄率" in chunks[0]["text"]
