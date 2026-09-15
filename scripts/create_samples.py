"""Create original, shareable text fixtures. No generated exam answers are bundled."""
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
TOPICS = [
    ("关系模型与 SQL", [
        "关系数据库以二维表组织数据。表中的一行称为元组，一列称为属性。主键用于唯一标识每一行，主键值不能重复，也不能为 NULL。",
        "外键引用另一张表的候选键，用于维护参照完整性。外键约束能防止产生无对应父记录的引用，但不自动代替所有业务规则。",
        "SELECT 用于查询，WHERE 对行进行筛选，GROUP BY 将行分组，HAVING 对分组结果筛选，ORDER BY 对输出排序。WHERE 不能直接代替 HAVING 对聚合结果筛选。",
        "COUNT(*) 统计行数，COUNT(列名) 只统计该列非 NULL 的行数。SUM 和 AVG 用于数值聚合，NULL 通常不参与聚合计算。",
        "内连接仅返回两表满足连接条件的行；左外连接保留左表所有行，右侧未匹配的属性以 NULL 填充。连接条件遗漏可能产生笛卡尔积。",
    ]),
    ("索引与查询成本", [
        "索引用额外存储维护可快速检索的数据结构，能减少满足条件的数据访问量，但增加插入、更新与删除的维护开销。不是所有查询都能从索引受益。",
        "B+树的内部节点主要保存键与子节点指针，记录或记录指针在叶子层。较高扇出能降低树高，叶子节点有序相连便于范围查询。",
        "复合索引按多个列的键顺序组织，通常遵循最左前缀原则。索引 (a,b,c) 一般能用于以 a 或 (a,b) 开头的等值筛选，不能据此保证仅筛选 b 的查询同样高效。",
        "选择性表示筛选后保留记录的比例。对于返回大量记录的查询，优化器可能选择全表扫描。索引存在不意味着优化器一定使用它。",
        "在简化 I/O 成本模型下，全表扫描页数为记录数乘每条记录字节数，再除以每页可用字节数并向上取整。假设忽略页头和跨页开销，1000 条、每条 80 字节、每页 4000 字节需 20 页。",
        "若 B+树查找路径需要 3 次页读取，再读取目标数据页 1 次，则一次点查询简化为 4 次页读取。这个估算假设缓存为空且每层访问一个页，不代表所有实际查询成本。",
    ]),
    ("事务与恢复", [
        "事务是一个逻辑工作单元，通常由多个操作组成。原子性要求事务中的操作全部完成，或在失败时回滚，避免只执行其中一部分。",
        "一致性要求事务将数据库从一个满足约束的状态转变为另一个满足约束的状态。隔离性用于控制并发事务之间的相互影响。持久性要求已提交结果在故障后仍能保存。",
        "转账案例：账户 A 向 B 转账 100 元，包含 A 减少 100 与 B 增加 100。两步必须同时成功或同时撤销，才能避免仅扣款不入账的中间结果。",
        "事务提交表示确认操作结果，回滚表示撤销尚未提交的修改。预写日志原则要求相关日志先持久化，再把修改后的数据页写回持久存储。",
        "检查点有助于缩小崩溃恢复时需要处理的日志范围。恢复过程可能需要重做已提交事务的修改，以及撤销未提交事务的修改，具体策略取决于实现。",
    ]),
    ("隔离级别与并发", [
        "脏读是读取到另一个事务尚未提交的数据。如果该事务随后回滚，读者就曾依赖一个无效结果。",
        "不可重复读是同一事务两次读取同一行时，因其他事务已提交的更新而得到不同内容。幻读通常涉及同一查询条件匹配的行集合变化。",
        "SQL 标准的读已提交隔离级别避免脏读，但可能发生不可重复读。可串行化要求并发执行效果等价于某个串行顺序，通常带来更高的并发控制成本。",
        "共享锁允许多个事务并发读取同一对象，排他锁用于独占修改。具体锁兼容性依赖锁粒度和数据库实现，不能简单认为读操作永远不加锁。",
        "死锁可能发生在多个事务循环等待对方持有的锁时。统一资源访问顺序、缩短事务、检测死锁后回滚一个事务都是常见处理方式。",
    ]),
]

def create(directory: Path):
    directory.mkdir(exist_ok=True, parents=True)
    doc = Document()
    doc.add_heading('数据库原理 · ExamPilot 原创演示讲义', 0)
    doc.add_paragraph('本讲义由项目编写，用于演示资料解析、知识提取与有据出题，可自由用于本项目展示。内容采用教学简化条件。')
    for title, paragraphs in TOPICS:
        doc.add_heading(title, 1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
    doc.save(directory / '数据库原理-原创演示讲义.docx')
    deck = Presentation()
    for title, paragraphs in TOPICS:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = '\n'.join(paragraphs[:3])
    deck.save(directory / '数据库原理-文字课件.pptx')

if __name__ == '__main__':
    create(ROOT / 'samples')
